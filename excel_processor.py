import os
import re
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.chart import LineChart, BarChart, DoughnutChart, Reference
from openpyxl.chart.shapes import GraphicalProperties
from openpyxl.formatting.rule import CellIsRule

def clean_location_name(full_name, common_prefix=""):
    """
    Remove prefixo comum de obra/empresa se presente, deixando o nome limpo do local.
    Ex: 'UFV NORONHA VERDE - USO E CONSUMO' -> 'USO E CONSUMO'
    """
    if not full_name:
        return ""
    name_str = str(full_name).strip()
    if common_prefix and name_str.startswith(common_prefix):
        name_str = name_str[len(common_prefix):].lstrip(" -").strip()
    elif " - " in name_str:
        parts = name_str.split(" - ", 1)
        name_str = parts[1].strip()
    return name_str.upper()

def format_currency_br(val):
    try:
        val_f = float(val)
        return f"R$ {val_f:,.2f}".replace(",", "X").replace(".", ",").replace("X", ".")
    except:
        return "R$ 0,00"

def process_inventory_excel(input_path, output_dir=None):
    """
    Processa a planilha bruta gerada pelo sistema e produz o arquivo
    completamente enriquecido com Curva ABC e Dashboard.
    Retorna o caminho do arquivo gerado e estatísticas completas para a interface.
    """
    if not os.path.exists(input_path):
        raise FileNotFoundError(f"Arquivo não encontrado: {input_path}")

    # 1. Carregar workbook bruto com dados e com fórmulas
    wb_data = openpyxl.load_workbook(input_path, data_only=True)
    wb_formulas = openpyxl.load_workbook(input_path, data_only=False)

    orig_sheetnames = wb_data.sheetnames
    if not orig_sheetnames:
        raise ValueError("O arquivo Excel enviado está vazio ou não possui abas.")

    # Extrair data a partir do nome do arquivo ou usar data padrão
    date_str = "27/09/2026"
    filename = os.path.basename(input_path)
    m_date = re.search(r'(\d{2})(\d{2})(\d{4})', filename)
    if m_date:
        d, m, y = m_date.groups()
        date_str = f"{d}/{m}/{y}"

    # Identificar o prefixo do projeto a partir do nome das abas ou locais
    project_name = "UFV NORONHA VERDE"
    common_prefix = ""
    for sname in orig_sheetnames:
        ws = wb_data[sname]
        for r in range(2, min(10, ws.max_row)):
            val = ws.cell(r, 2).value
            if val and " - " in str(val):
                possible_proj = str(val).split(" - ")[0].strip()
                if possible_proj:
                    project_name = possible_proj
                    common_prefix = possible_proj + " -"
                    break
        if common_prefix:
            break

    # Mapear abas para nomes padronizados limpos
    sheet_renames = {}
    for idx, sname in enumerate(orig_sheetnames):
        sl = sname.lower()
        if 'consumo' in sl or sname == 'Sheet' or (idx == 0 and len(orig_sheetnames) >= 3):
            clean = 'Uso e Consumo'
        elif 'fat' in sl or 'terceiro' in sl or (idx == 1 and len(orig_sheetnames) >= 3):
            clean = 'Fat Dir Terceiro'
        elif 'imobilizado' in sl or (idx == 2 and len(orig_sheetnames) >= 3):
            clean = 'Imobilizado'
        else:
            clean = sname[:31]
        sheet_renames[sname] = clean

    # Criar um novo workbook onde montaremos a estrutura final
    wb_out = openpyxl.Workbook()
    default_sheet = wb_out.active

    # Criar abas na ordem exata solicitada:
    # 1. Dashboard
    # 2. Curva ABC
    # 3. Análises
    # 4. Base
    # 5..N. Abas de dados renomeadas
    ws_dash = wb_out.create_sheet('Dashboard')
    ws_abc = wb_out.create_sheet('Curva ABC')
    ws_analises = wb_out.create_sheet('Análises')
    ws_base = wb_out.create_sheet('Base')

    wb_out.remove(default_sheet)

    # Copiar as abas de dados originais para wb_out mantendo dados e fórmulas
    data_sheets_info = []

    for orig_name in orig_sheetnames:
        clean_name = sheet_renames[orig_name]
        ws_dest = wb_out.create_sheet(clean_name)
        ws_src_f = wb_formulas[orig_name]
        ws_src_d = wb_data[orig_name]

        max_r = ws_src_f.max_row
        max_c = ws_src_f.max_column

        # Determinar linha de TOTAL
        total_row = None
        for r in range(max_r, max(1, max_r - 5), -1):
            val_g_f = str(ws_src_f.cell(r, 7).value or "")
            val_f_d = str(ws_src_d.cell(r, 6).value or "").upper()
            if "SUM(" in val_g_f.upper() or val_f_d == "TOTAL":
                total_row = r
                break

        data_end_row = (total_row - 1) if total_row else max_r

        # Copiar células
        for r in range(1, max_r + 1):
            for c in range(1, max_c + 1):
                c_src = ws_src_f.cell(r, c)
                c_dest = ws_dest.cell(r, c, c_src.value)
                if c_src.number_format:
                    c_dest.number_format = c_src.number_format

        # Garantir que a linha de total tenha 'TOTAL' na coluna F (6)
        if total_row:
            ws_dest.cell(total_row, 6, 'TOTAL')
            ws_dest.cell(total_row, 6).font = Font(name="Arial", size=10, bold=True)
            ws_dest.cell(total_row, 7).font = Font(name="Arial", size=10, bold=True)

        data_sheets_info.append({
            'orig_name': orig_name,
            'clean_name': clean_name,
            'data_start': 2,
            'data_end': data_end_row,
            'total_row': total_row
        })

    # =========================================================================
    # LEITURA E CONSOLIDAÇÃO DOS DADOS
    # =========================================================================
    items_dict = {}
    loc_totals = {}
    natureza_totals = {}
    total_base_rows = 0
    total_inventory_value = 0.0

    for s_info in data_sheets_info:
        orig_name = s_info['orig_name']
        ws_d = wb_data[orig_name]
        for r in range(s_info['data_start'], s_info['data_end'] + 1):
            codloc = str(ws_d.cell(r, 1).value or "").strip()
            loc_nome = str(ws_d.cell(r, 2).value or "").strip()
            idprd = ws_d.cell(r, 3).value
            codigoprd = str(ws_d.cell(r, 4).value or "").strip()
            nomefantasia = str(ws_d.cell(r, 5).value or "").strip()
            saldo_fin = ws_d.cell(r, 7).value or 0.0
            natureza = str(ws_d.cell(r, 13).value or "").strip()

            if idprd is None:
                continue

            try:
                val_num = float(saldo_fin)
            except (ValueError, TypeError):
                val_num = 0.0

            total_base_rows += 1
            total_inventory_value += val_num

            if idprd not in items_dict:
                items_dict[idprd] = {
                    'idprd': idprd,
                    'codigoprd': codigoprd,
                    'nome': nomefantasia,
                    'natureza': natureza,
                    'total_val': 0.0,
                    'locations': set()
                }
            items_dict[idprd]['total_val'] += val_num
            items_dict[idprd]['locations'].add(codloc)

            # Locais
            if codloc not in loc_totals:
                loc_clean_nome = clean_location_name(loc_nome, common_prefix)
                loc_totals[codloc] = {
                    'codloc': codloc,
                    'nome': loc_clean_nome or loc_nome,
                    'total_val': 0.0,
                    'count': 0
                }
            loc_totals[codloc]['total_val'] += val_num
            loc_totals[codloc]['count'] += 1

            # Natureza
            if natureza:
                if natureza not in natureza_totals:
                    natureza_totals[natureza] = {'total_val': 0.0, 'count': 0}
                natureza_totals[natureza]['total_val'] += val_num
                natureza_totals[natureza]['count'] += 1

    # Ordenar itens por total_val decrescente
    sorted_items = sorted(items_dict.values(), key=lambda x: x['total_val'], reverse=True)
    num_products = len(sorted_items)

    # Ordenar locais por total_val decrescente
    sorted_locations = sorted(loc_totals.values(), key=lambda x: x['total_val'], reverse=True)

    # Ordenar naturezas por total_val decrescente
    sorted_naturezas = sorted(natureza_totals.items(), key=lambda x: x[1]['total_val'], reverse=True)

    # Calcular corte ABC em memória para o frontend
    accum_val = 0.0
    count_a = 0
    val_a = 0.0
    count_b = 0
    val_b = 0.0
    count_c = 0
    val_c = 0.0

    for itm in sorted_items:
        v = itm['total_val']
        accum_val += v
        pct_accum = (accum_val / total_inventory_value) if total_inventory_value > 0 else 0
        if pct_accum <= 0.80:
            count_a += 1
            val_a += v
        elif pct_accum <= 0.95:
            count_b += 1
            val_b += v
        else:
            count_c += 1
            val_c += v

    # =========================================================================
    # DEFINIÇÃO DE LINHAS DAS ABAS
    # =========================================================================
    last_base_row = 1 + total_base_rows
    first_abc_row = 8
    last_abc_row = first_abc_row + num_products - 1
    total_abc_row = last_abc_row + 2
    conf_abc_row = last_abc_row + 3

    # =========================================================================
    # 1. ABA BASE
    # =========================================================================
    ws_base.column_dimensions['A'].width = 18.0
    ws_base.column_dimensions['B'].width = 9.0
    ws_base.column_dimensions['C'].width = 40.0
    ws_base.column_dimensions['D'].width = 11.0
    ws_base.column_dimensions['E'].width = 13.0
    ws_base.column_dimensions['F'].width = 70.0
    ws_base.column_dimensions['G'].width = 8.0
    ws_base.column_dimensions['H'].width = 13.0
    ws_base.column_dimensions['I'].width = 19.0
    ws_base.column_dimensions['J'].width = 46.0
    ws_base.column_dimensions['K'].width = 11.0

    base_headers = [
        'Aba origem', 'CODLOC', 'Local', 'IDPRD', 'Código',
        'Descrição', 'Unid.', 'Saldo físico', 'Saldo financeiro (R$)',
        'Natureza', 'Classe ABC'
    ]
    for c_idx, h_text in enumerate(base_headers, 1):
        cell = ws_base.cell(1, c_idx, h_text)
        cell.font = Font(name="Arial", size=10, bold=True, color="FFFFFF")
        cell.fill = PatternFill(start_color="1F3864", end_color="1F3864", fill_type="solid")
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)

    curr_base_r = 2
    for s_info in data_sheets_info:
        clean_name = s_info['clean_name']
        for r_src in range(s_info['data_start'], s_info['data_end'] + 1):
            ws_base.cell(curr_base_r, 1, clean_name)
            ws_base.cell(curr_base_r, 2, f"='{clean_name}'!A{r_src}")
            ws_base.cell(curr_base_r, 3, f"='{clean_name}'!B{r_src}")
            ws_base.cell(curr_base_r, 4, f"='{clean_name}'!C{r_src}")
            ws_base.cell(curr_base_r, 5, f"='{clean_name}'!D{r_src}")
            ws_base.cell(curr_base_r, 6, f"='{clean_name}'!E{r_src}")
            ws_base.cell(curr_base_r, 7, f"='{clean_name}'!J{r_src}")
            
            c_f = ws_base.cell(curr_base_r, 8, f"='{clean_name}'!F{r_src}")
            c_f.number_format = '#,##0.00'
            
            c_g = ws_base.cell(curr_base_r, 9, f"='{clean_name}'!G{r_src}")
            c_g.number_format = '"R$ "#,##0.00'
            
            ws_base.cell(curr_base_r, 10, f"='{clean_name}'!M{r_src}")
            
            c_k = ws_base.cell(curr_base_r, 11, f"=INDEX('Curva ABC'!$K${first_abc_row}:$K${last_abc_row},MATCH(D{curr_base_r},'Curva ABC'!$B${first_abc_row}:$B${last_abc_row},0))")
            c_k.alignment = Alignment(horizontal="center")
            
            for col in range(1, 12):
                ws_base.cell(curr_base_r, col).font = Font(name="Arial", size=9)

            curr_base_r += 1

    # =========================================================================
    # 2. ABA CURVA ABC
    # =========================================================================
    ws_abc.column_dimensions['A'].width = 7.0
    ws_abc.column_dimensions['B'].width = 11.0
    ws_abc.column_dimensions['C'].width = 13.0
    ws_abc.column_dimensions['D'].width = 68.0
    ws_abc.column_dimensions['E'].width = 44.0
    ws_abc.column_dimensions['F'].width = 9.0
    ws_abc.column_dimensions['G'].width = 20.0
    ws_abc.column_dimensions['H'].width = 10.0
    ws_abc.column_dimensions['I'].width = 12.0
    ws_abc.column_dimensions['J'].width = 12.0
    ws_abc.column_dimensions['K'].width = 8.0

    ws_abc.cell(1, 1, 'CURVA ABC DE ESTOQUE – CONSOLIDADA POR PRODUTO').font = Font(name="Arial", size=14, bold=True, color="1F3864")
    ws_abc.cell(2, 1, f"Valor = soma do saldo financeiro do produto em todos os locais (aba Base). Ordem por valor decrescente (ordenada em {date_str}).").font = Font(name="Arial", size=9, color="595959")

    ws_abc.cell(4, 1, 'Limite classe A (% acumulado do valor)').font = Font(name="Arial", size=10, bold=True)
    c_lim_a = ws_abc.cell(4, 5, 0.80)
    c_lim_a.font = Font(name="Arial", size=10, bold=True, color="0000FF")
    c_lim_a.fill = PatternFill(start_color="FFFF00", end_color="FFFF00", fill_type="solid")
    c_lim_a.alignment = Alignment(horizontal="center")
    c_lim_a.number_format = '0%'

    ws_abc.cell(4, 7, '← células amarelas: edite para mudar o corte (A: acumulado ≤ limite A; B: ≤ limite B; C: restante)').font = Font(name="Arial", size=9, color="595959")

    ws_abc.cell(5, 1, 'Limite classe B (% acumulado do valor)').font = Font(name="Arial", size=10, bold=True)
    c_lim_b = ws_abc.cell(5, 5, 0.95)
    c_lim_b.font = Font(name="Arial", size=10, bold=True, color="0000FF")
    c_lim_b.fill = PatternFill(start_color="FFFF00", end_color="FFFF00", fill_type="solid")
    c_lim_b.alignment = Alignment(horizontal="center")
    c_lim_b.number_format = '0%'

    abc_headers = [
        'Rank', 'IDPRD', 'Código', 'Descrição', 'Natureza',
        'Nº locais', 'Valor em estoque (R$)', '% do valor',
        '% acumulado', '% acum. itens', 'Classe'
    ]
    thin_border = Border(
        left=Side(style='thin', color='D9D9D9'),
        right=Side(style='thin', color='D9D9D9'),
        top=Side(style='thin', color='D9D9D9'),
        bottom=Side(style='thin', color='D9D9D9')
    )
    for c_idx, h_text in enumerate(abc_headers, 1):
        cell = ws_abc.cell(7, c_idx, h_text)
        cell.font = Font(name="Arial", size=10, bold=True, color="FFFFFF")
        cell.fill = PatternFill(start_color="1F3864", end_color="1F3864", fill_type="solid")
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        cell.border = thin_border

    for i, itm in enumerate(sorted_items):
        r_abc = first_abc_row + i
        rank = i + 1
        
        c_a = ws_abc.cell(r_abc, 1, rank)
        c_a.alignment = Alignment(horizontal="center")
        
        c_b = ws_abc.cell(r_abc, 2, itm['idprd'])
        c_b.number_format = '0'
        
        ws_abc.cell(r_abc, 3, itm['codigoprd'])
        ws_abc.cell(r_abc, 4, itm['nome'])
        ws_abc.cell(r_abc, 5, itm['natureza'])
        
        c_f = ws_abc.cell(r_abc, 6, f"=COUNTIF(Base!$D$2:$D${last_base_row},B{r_abc})")
        c_f.alignment = Alignment(horizontal="center")
        
        c_g = ws_abc.cell(r_abc, 7, f"=SUMIF(Base!$D$2:$D${last_base_row},B{r_abc},Base!$I$2:$I${last_base_row})")
        c_g.number_format = '"R$ "#,##0.00'
        
        c_h = ws_abc.cell(r_abc, 8, f"=IF($G${total_abc_row}=0,0,G{r_abc}/$G${total_abc_row})")
        c_h.number_format = '0.00%'
        
        if rank == 1:
            c_i = ws_abc.cell(r_abc, 9, f"=H{r_abc}")
        else:
            c_i = ws_abc.cell(r_abc, 9, f"=I{r_abc-1}+H{r_abc}")
        c_i.number_format = '0.0%'
        
        c_j = ws_abc.cell(r_abc, 10, f"=A{r_abc}/{num_products}")
        c_j.number_format = '0.0%'
        
        c_k = ws_abc.cell(r_abc, 11, f'=IF(I{r_abc}<=$E$4,"A",IF(I{r_abc}<=$E$5,"B","C"))')
        c_k.alignment = Alignment(horizontal="center")

        for col in range(1, 12):
            ws_abc.cell(r_abc, col).font = Font(name="Arial", size=9)

    # Linha Total
    ws_abc.cell(total_abc_row, 6, 'TOTAL').font = Font(name="Arial", size=10, bold=True)
    ws_abc.cell(total_abc_row, 6).alignment = Alignment(horizontal="right")
    c_tot_g = ws_abc.cell(total_abc_row, 7, f"=SUM(G{first_abc_row}:G{last_abc_row})")
    c_tot_g.font = Font(name="Arial", size=10, bold=True)
    c_tot_g.number_format = '"R$ "#,##0.00'

    # Linha Conferência
    ws_abc.cell(conf_abc_row, 6, 'Conferência vs Base').font = Font(name="Arial", size=9, color="595959")
    ws_abc.cell(conf_abc_row, 6).alignment = Alignment(horizontal="right")
    c_conf_g = ws_abc.cell(conf_abc_row, 7, f"=G{total_abc_row}-SUM(Base!$I$2:$I${last_base_row})")
    c_conf_g.font = Font(name="Arial", size=9, color="595959")
    c_conf_g.number_format = '"R$ "#,##0.00'

    # Formatação condicional
    green_fill = PatternFill(start_color="C6EFCE", end_color="C6EFCE", fill_type="solid")
    green_font = Font(color="006100")
    yellow_fill = PatternFill(start_color="FFEB9C", end_color="FFEB9C", fill_type="solid")
    yellow_font = Font(color="9C6500")
    red_fill = PatternFill(start_color="F8CBAD", end_color="F8CBAD", fill_type="solid")
    red_font = Font(color="9C0006")

    abc_range = f"K{first_abc_row}:K{last_abc_row}"
    ws_abc.conditional_formatting.add(abc_range, CellIsRule(operator='equal', formula=['"A"'], fill=green_fill, font=green_font))
    ws_abc.conditional_formatting.add(abc_range, CellIsRule(operator='equal', formula=['"B"'], fill=yellow_fill, font=yellow_font))
    ws_abc.conditional_formatting.add(abc_range, CellIsRule(operator='equal', formula=['"C"'], fill=red_fill, font=red_font))

    # =========================================================================
    # 3. ABA ANÁLISES
    # =========================================================================
    ws_analises.column_dimensions['A'].width = 58.0
    ws_analises.column_dimensions['B'].width = 40.0
    ws_analises.column_dimensions['C'].width = 17.0
    ws_analises.column_dimensions['D'].width = 17.0
    ws_analises.column_dimensions['E'].width = 14.0
    ws_analises.column_dimensions['F'].width = 18.0

    ws_analises.cell(1, 1, 'TABELAS DE APOIO DO DASHBOARD (fórmulas – não editar)').font = Font(name="Arial", size=11, bold=True, color="1F3864")

    # Tabela 1: Resumo ABC (linhas 3 a 7)
    an_t1_headers = ['Classe', 'Nº itens', '% itens', 'Valor (R$)', '% valor']
    for ci, h in enumerate(an_t1_headers, 1):
        c = ws_analises.cell(3, ci, h)
        c.font = Font(name="Arial", size=10, bold=True, color="FFFFFF")
        c.fill = PatternFill(start_color="1F3864", end_color="1F3864", fill_type="solid")
        c.alignment = Alignment(horizontal="center")
        c.border = thin_border

    classes = ['A', 'B', 'C']
    for idx, cl in enumerate(classes):
        r = 4 + idx
        ws_analises.cell(r, 1, cl).alignment = Alignment(horizontal="center")
        
        c_b = ws_analises.cell(r, 2, f"=COUNTIF('Curva ABC'!$K${first_abc_row}:$K${last_abc_row},A{r})")
        c_b.number_format = '#,##0'
        
        c_c = ws_analises.cell(r, 3, f"=B{r}/$B$7")
        c_c.number_format = '0.0%'
        
        c_d = ws_analises.cell(r, 4, f"=SUMIF('Curva ABC'!$K${first_abc_row}:$K${last_abc_row},A{r},'Curva ABC'!$G${first_abc_row}:$G${last_abc_row})")
        c_d.number_format = '"R$ "#,##0.00'
        
        c_e = ws_analises.cell(r, 5, f"=D{r}/$D$7")
        c_e.number_format = '0.0%'

    # Total ABC (Linha 7)
    ws_analises.cell(7, 1, 'Total').font = Font(name="Arial", size=10, bold=True)
    c_b7 = ws_analises.cell(7, 2, "=SUM(B4:B6)")
    c_b7.number_format = '#,##0'
    c_b7.font = Font(name="Arial", size=10, bold=True)
    
    c_c7 = ws_analises.cell(7, 3, "=SUM(C4:C6)")
    c_c7.number_format = '0.0%'
    c_c7.font = Font(name="Arial", size=10, bold=True)
    
    c_d7 = ws_analises.cell(7, 4, "=SUM(D4:D6)")
    c_d7.number_format = '"R$ "#,##0.00'
    c_d7.font = Font(name="Arial", size=10, bold=True)
    
    c_e7 = ws_analises.cell(7, 5, "=SUM(E4:E6)")
    c_e7.number_format = '0.0%'
    c_e7.font = Font(name="Arial", size=10, bold=True)

    # Tabela 2: Resumo por Local (linhas 10 a 16)
    an_t2_headers = ['CODLOC', 'Local', 'Nº linhas', 'Valor (R$)', '% valor']
    for ci, h in enumerate(an_t2_headers, 1):
        c = ws_analises.cell(10, ci, h)
        c.font = Font(name="Arial", size=10, bold=True, color="FFFFFF")
        c.fill = PatternFill(start_color="1F3864", end_color="1F3864", fill_type="solid")
        c.alignment = Alignment(horizontal="center")
        c.border = thin_border

    num_locs = len(sorted_locations)
    first_loc_r = 11
    last_loc_r = first_loc_r + num_locs - 1
    tot_loc_r = last_loc_r + 1

    for idx, loc in enumerate(sorted_locations):
        r = first_loc_r + idx
        ws_analises.cell(r, 1, loc['codloc']).alignment = Alignment(horizontal="center")
        ws_analises.cell(r, 2, loc['nome'])
        
        c_c = ws_analises.cell(r, 3, f"=COUNTIF(Base!$B$2:$B${last_base_row},A{r})")
        c_c.number_format = '#,##0'
        
        c_d = ws_analises.cell(r, 4, f"=SUMIF(Base!$B$2:$B${last_base_row},A{r},Base!$I$2:$I${last_base_row})")
        c_d.number_format = '"R$ "#,##0.00'
        
        c_e = ws_analises.cell(r, 5, f"=D{r}/$D${tot_loc_r}")
        c_e.number_format = '0.0%'

    # Total Local
    ws_analises.cell(tot_loc_r, 1, 'Total').font = Font(name="Arial", size=10, bold=True)
    c_cl = ws_analises.cell(tot_loc_r, 3, f"=SUM(C{first_loc_r}:C{last_loc_r})")
    c_cl.number_format = '#,##0'
    c_cl.font = Font(name="Arial", size=10, bold=True)
    
    c_dl = ws_analises.cell(tot_loc_r, 4, f"=SUM(D{first_loc_r}:D{last_loc_r})")
    c_dl.number_format = '"R$ "#,##0.00'
    c_dl.font = Font(name="Arial", size=10, bold=True)
    
    c_el = ws_analises.cell(tot_loc_r, 5, f"=SUM(E{first_loc_r}:E{last_loc_r})")
    c_el.number_format = '0.0%'
    c_el.font = Font(name="Arial", size=10, bold=True)

    # Tabela 3: Matriz Local x Classe
    mat_start_r = tot_loc_r + 3
    an_t3_headers = ['CODLOC', 'Local', 'Classe A (R$)', 'Classe B (R$)', 'Classe C (R$)', 'Total (R$)']
    for ci, h in enumerate(an_t3_headers, 1):
        c = ws_analises.cell(mat_start_r, ci, h)
        c.font = Font(name="Arial", size=10, bold=True, color="FFFFFF")
        c.fill = PatternFill(start_color="1F3864", end_color="1F3864", fill_type="solid")
        c.alignment = Alignment(horizontal="center")
        c.border = thin_border

    first_mat_r = mat_start_r + 1
    for idx, loc in enumerate(sorted_locations):
        r = first_mat_r + idx
        loc_ref_r = first_loc_r + idx
        ws_analises.cell(r, 1, loc['codloc']).alignment = Alignment(horizontal="center")
        ws_analises.cell(r, 2, f"=B{loc_ref_r}")
        
        c_c = ws_analises.cell(r, 3, f'=SUMIFS(Base!$I$2:$I${last_base_row},Base!$B$2:$B${last_base_row},$A{r},Base!$K$2:$K${last_base_row},"A")')
        c_c.number_format = '"R$ "#,##0.00'
        
        c_d = ws_analises.cell(r, 4, f'=SUMIFS(Base!$I$2:$I${last_base_row},Base!$B$2:$B${last_base_row},$A{r},Base!$K$2:$K${last_base_row},"B")')
        c_d.number_format = '"R$ "#,##0.00'
        
        c_e = ws_analises.cell(r, 5, f'=SUMIFS(Base!$I$2:$I${last_base_row},Base!$B$2:$B${last_base_row},$A{r},Base!$K$2:$K${last_base_row},"C")')
        c_e.number_format = '"R$ "#,##0.00'
        
        c_f = ws_analises.cell(r, 6, f"=SUM(C{r}:E{r})")
        c_f.number_format = '"R$ "#,##0.00'

    # Tabela 4: Top 10 Naturezas por valor
    nat_start_r = 28
    an_t4_headers = ['Natureza (top 10)', 'Nº itens', 'Valor (R$)', '% valor']
    for ci, h in enumerate(an_t4_headers, 1):
        c = ws_analises.cell(nat_start_r, ci, h)
        c.font = Font(name="Arial", size=10, bold=True, color="FFFFFF")
        c.fill = PatternFill(start_color="1F3864", end_color="1F3864", fill_type="solid")
        c.alignment = Alignment(horizontal="center")
        c.border = thin_border

    top_10_nat = sorted_naturezas[:10]
    for idx, (nat_name, nat_data) in enumerate(top_10_nat):
        r = nat_start_r + 1 + idx
        ws_analises.cell(r, 1, nat_name)
        
        c_b = ws_analises.cell(r, 2, f"=COUNTIF('Curva ABC'!$E${first_abc_row}:$E${last_abc_row},A{r})")
        c_b.number_format = '#,##0'
        
        c_c = ws_analises.cell(r, 3, f"=SUMIF('Curva ABC'!$E${first_abc_row}:$E${last_abc_row},A{r},'Curva ABC'!$G${first_abc_row}:$G${last_abc_row})")
        c_c.number_format = '"R$ "#,##0.00'
        
        c_d = ws_analises.cell(r, 4, f"=C{r}/'Curva ABC'!$G${total_abc_row}")
        c_d.number_format = '0.0%'

    demais_nat_r = nat_start_r + 1 + len(top_10_nat)
    ws_analises.cell(demais_nat_r, 1, 'Demais naturezas').font = Font(name="Arial", size=10, italic=True)
    c_cd = ws_analises.cell(demais_nat_r, 3, f"='Curva ABC'!$G${total_abc_row}-SUM(C29:C{demais_nat_r-1})")
    c_cd.number_format = '"R$ "#,##0.00'
    c_dd = ws_analises.cell(demais_nat_r, 4, f"=C{demais_nat_r}/'Curva ABC'!$G${total_abc_row}")
    c_dd.number_format = '0.0%'

    # Tabela 5: Top 10 itens
    t5_start_r = 43
    an_t5_headers = ['Top 10 itens', 'Valor (R$)', '% valor']
    for ci, h in enumerate(an_t5_headers, 1):
        c = ws_analises.cell(t5_start_r, ci, h)
        c.font = Font(name="Arial", size=10, bold=True, color="FFFFFF")
        c.fill = PatternFill(start_color="1F3864", end_color="1F3864", fill_type="solid")
        c.alignment = Alignment(horizontal="center")
        c.border = thin_border

    for idx in range(min(10, num_products)):
        r = t5_start_r + 1 + idx
        abc_r = first_abc_row + idx
        ws_analises.cell(r, 1, f"='Curva ABC'!C{abc_r}&\" - \"&LEFT('Curva ABC'!D{abc_r},38)")
        c_b = ws_analises.cell(r, 2, f"='Curva ABC'!G{abc_r}")
        c_b.number_format = '"R$ "#,##0.00'
        c_c = ws_analises.cell(r, 3, f"='Curva ABC'!H{abc_r}")
        c_c.number_format = '0.00%'

    # Tabela 6: Top 20 itens por rank
    t6_start_r = 57
    an_t6_headers = ['Rank', 'Valor (R$)', '% acumulado']
    for ci, h in enumerate(an_t6_headers, 1):
        c = ws_analises.cell(t6_start_r, ci, h)
        c.font = Font(name="Arial", size=10, bold=True, color="FFFFFF")
        c.fill = PatternFill(start_color="1F3864", end_color="1F3864", fill_type="solid")
        c.alignment = Alignment(horizontal="center")
        c.border = thin_border

    for idx in range(min(20, num_products)):
        r = t6_start_r + 1 + idx
        abc_r = first_abc_row + idx
        ws_analises.cell(r, 1, f"='Curva ABC'!A{abc_r}").alignment = Alignment(horizontal="center")
        c_b = ws_analises.cell(r, 2, f"='Curva ABC'!G{abc_r}")
        c_b.number_format = '"R$ "#,##0.00'
        c_c = ws_analises.cell(r, 3, f"='Curva ABC'!I{abc_r}")
        c_c.number_format = '0.0%'

    # =========================================================================
    # 4. ABA DASHBOARD
    # =========================================================================
    ws_dash.column_dimensions['A'].width = 9.51

    # Cabeçalho Principal (A1:Z2)
    ws_dash.merge_cells('A1:Z2')
    c_title = ws_dash['A1']
    c_title.value = f"INVENTÁRIO – {project_name}  |  DASHBOARD & CURVA ABC"
    c_title.font = Font(name="Arial", size=14, bold=True, color="FFFFFF")
    c_title.fill = PatternFill(start_color="1F3864", end_color="1F3864", fill_type="solid")
    c_title.alignment = Alignment(horizontal="center", vertical="center")

    # Subtítulo (A3:Z3)
    ws_dash.merge_cells('A3:Z3')
    c_sub = ws_dash['A3']
    c_sub.value = f"Posição de estoque em {date_str}  •  fonte: abas originais do inventário  •  valores em R$"
    c_sub.font = Font(name="Arial", size=9, italic=False, color="595959")
    c_sub.alignment = Alignment(horizontal="center", vertical="center")

    # CARDS DE KPI (Linha 5 e Linha 6-7)
    kpis = [
        ('A5:D5', 'A6:D7', 'VALOR TOTAL EM ESTOQUE', f"='Curva ABC'!G{total_abc_row}", '"R$ "#,##0.00'),
        ('E5:H5', 'E6:H7', 'PRODUTOS DISTINTOS', f"=COUNT('Curva ABC'!$B${first_abc_row}:$B${last_abc_row})", '#,##0'),
        ('I5:L5', 'I6:L7', 'LOCAIS DE ESTOQUE', f"=COUNTA(Análises!A{first_loc_r}:A{last_loc_r})", '0'),
        ('M5:P5', 'M6:P7', 'ITENS CLASSE A', "=Análises!B4", '#,##0'),
        ('Q5:T5', 'Q6:T7', '% DO VALOR EM CLASSE A', "=Análises!E4", '0.0%'),
        ('U5:Z5', 'U6:Z7', 'MAIOR ITEM (% DO TOTAL)', f"='Curva ABC'!H{first_abc_row}", '0.0%')
    ]

    for t_range, v_range, title_text, formula_val, num_fmt in kpis:
        ws_dash.merge_cells(t_range)
        t_cell = ws_dash[t_range.split(':')[0]]
        t_cell.value = title_text
        t_cell.font = Font(name="Arial", size=9, bold=True, color="FFFFFF")
        t_cell.fill = PatternFill(start_color="2F5597", end_color="2F5597", fill_type="solid")
        t_cell.alignment = Alignment(horizontal="center", vertical="center")

        ws_dash.merge_cells(v_range)
        v_cell = ws_dash[v_range.split(':')[0]]
        v_cell.value = formula_val
        v_cell.font = Font(name="Arial", size=16, bold=True, color="1F3864")
        v_cell.fill = PatternFill(start_color="DDEBF7", end_color="DDEBF7", fill_type="solid")
        v_cell.alignment = Alignment(horizontal="center", vertical="center")
        v_cell.number_format = num_fmt

    # Seção Resumo ABC (Linhas 9 a 14)
    ws_dash.merge_cells('A9:Z9')
    c_sec = ws_dash['A9']
    c_sec.value = "Resumo por classe ABC"
    c_sec.font = Font(name="Arial", size=11, bold=True, color="1F3864")
    c_sec.alignment = Alignment(horizontal="left", vertical="center")

    resumo_cols = [
        ('A10:B10', 'Classe'),
        ('C10:D10', 'Nº itens'),
        ('E10:F10', '% itens'),
        ('G10:I10', 'Valor (R$)'),
        ('J10:K10', '% valor')
    ]
    for r_range, h_text in resumo_cols:
        ws_dash.merge_cells(r_range)
        cell = ws_dash[r_range.split(':')[0]]
        cell.value = h_text
        cell.font = Font(name="Arial", size=10, bold=True, color="FFFFFF")
        cell.fill = PatternFill(start_color="1F3864", end_color="1F3864", fill_type="solid")
        cell.alignment = Alignment(horizontal="center", vertical="center")

    ws_dash.merge_cells('M10:Z14')
    c_info = ws_dash['M10']
    c_info.value = "Como ler: poucos itens (A) concentram a maior parte do valor; priorize contagem, segurança e reposição neles."
    c_info.font = Font(name="Arial", size=9, color="595959")
    c_info.alignment = Alignment(horizontal="left", vertical="top", wrap_text=True)

    res_rows_data = [
        (11, 4, "C6EFCE", False),
        (12, 5, "FFEB9C", False),
        (13, 6, "F8CBAD", False),
        (14, 7, "DDEBF7", True),
    ]
    for dash_r, an_r, fill_color, is_bold in res_rows_data:
        cells_map = [
            (f"A{dash_r}:B{dash_r}", f"=Análises!A{an_r}", None, "center"),
            (f"C{dash_r}:D{dash_r}", f"=Análises!B{an_r}", '#,##0', "center"),
            (f"E{dash_r}:F{dash_r}", f"=Análises!C{an_r}", '0.0%', "center"),
            (f"G{dash_r}:I{dash_r}", f"=Análises!D{an_r}", '"R$ "#,##0.00', "right"),
            (f"J{dash_r}:K{dash_r}", f"=Análises!E{an_r}", '0.0%', "center")
        ]
        for c_range, formula_val, num_fmt, align_h in cells_map:
            ws_dash.merge_cells(c_range)
            c = ws_dash[c_range.split(':')[0]]
            c.value = formula_val
            c.font = Font(name="Arial", size=10, bold=is_bold)
            c.fill = PatternFill(start_color=fill_color, end_color=fill_color, fill_type="solid")
            c.alignment = Alignment(horizontal=align_h, vertical="center")
            if num_fmt:
                c.number_format = num_fmt

            start_col, start_row = c_range.split(':')[0][0], int(c_range.split(':')[0][1:])
            end_col, end_row = c_range.split(':')[1][0], int(c_range.split(':')[1][1:])
            for col_l in [start_col, end_col]:
                ws_dash[f"{col_l}{start_row}"].border = thin_border

    # =========================================================================
    # 5. OS 7 GRÁFICOS NO DASHBOARD
    # =========================================================================
    def apply_series_color(series, hex_color):
        sPr = GraphicalProperties()
        sPr.solidFill = hex_color
        series.graphicalProperties = sPr

    # Chart 0: LineChart - Curva ABC – % acumulado do valor
    chart0 = LineChart()
    chart0.title = "Curva ABC – % acumulado do valor"
    chart0.style = 13
    chart0.legend = None
    chart0.width = 16.5
    chart0.height = 7.5
    data_ref0 = Reference(ws_abc, min_col=9, min_row=first_abc_row, max_row=last_abc_row)
    cats_ref0 = Reference(ws_abc, min_col=1, min_row=first_abc_row, max_row=last_abc_row)
    chart0.add_data(data_ref0, titles_from_data=False)
    chart0.set_categories(cats_ref0)
    if chart0.series:
        apply_series_color(chart0.series[0], "C00000")
    ws_dash.add_chart(chart0, "A16")

    # Chart 1: BarChart (col) - Pareto – 20 maiores itens
    chart1 = BarChart()
    chart1.type = "col"
    chart1.grouping = "clustered"
    chart1.title = "Pareto – 20 maiores itens"
    chart1.legend = None
    chart1.width = 16.5
    chart1.height = 7.5
    data_ref1 = Reference(ws_analises, min_col=2, min_row=58, max_row=77)
    cats_ref1 = Reference(ws_analises, min_col=1, min_row=58, max_row=77)
    chart1.add_data(data_ref1, titles_from_data=False)
    chart1.set_categories(cats_ref1)
    if chart1.series:
        apply_series_color(chart1.series[0], "2F5597")
    ws_dash.add_chart(chart1, "N16")

    # Chart 2: DoughnutChart - Valor por classe ABC
    chart2 = DoughnutChart()
    chart2.title = "Valor por classe ABC"
    chart2.holeSize = 50
    chart2.width = 8.5
    chart2.height = 7.5
    data_ref2 = Reference(ws_analises, min_col=4, min_row=4, max_row=6)
    cats_ref2 = Reference(ws_analises, min_col=1, min_row=4, max_row=6)
    chart2.add_data(data_ref2, titles_from_data=False)
    chart2.set_categories(cats_ref2)
    ws_dash.add_chart(chart2, "A35")

    # Chart 3: BarChart (col) - Nº de itens por classe
    chart3 = BarChart()
    chart3.type = "col"
    chart3.grouping = "clustered"
    chart3.title = "Nº de itens por classe"
    chart3.legend = None
    chart3.width = 8.5
    chart3.height = 7.5
    data_ref3 = Reference(ws_analises, min_col=2, min_row=4, max_row=6)
    cats_ref3 = Reference(ws_analises, min_col=1, min_row=4, max_row=6)
    chart3.add_data(data_ref3, titles_from_data=False)
    chart3.set_categories(cats_ref3)
    if chart3.series:
        apply_series_color(chart3.series[0], "5B9BD5")
    ws_dash.add_chart(chart3, "I35")

    # Chart 4: BarChart (bar horizontal) - Valor por local de estoque
    chart4 = BarChart()
    chart4.type = "bar"
    chart4.grouping = "clustered"
    chart4.title = "Valor por local de estoque"
    chart4.legend = None
    chart4.width = 16.0
    chart4.height = 7.5
    data_ref4 = Reference(ws_analises, min_col=4, min_row=first_loc_r, max_row=last_loc_r)
    cats_ref4 = Reference(ws_analises, min_col=2, min_row=first_loc_r, max_row=last_loc_r)
    chart4.add_data(data_ref4, titles_from_data=False)
    chart4.set_categories(cats_ref4)
    if chart4.series:
        apply_series_color(chart4.series[0], "2F5597")
    ws_dash.add_chart(chart4, "Q35")

    # Chart 5: BarChart (bar horizontal) - Top 10 naturezas por valor
    chart5 = BarChart()
    chart5.type = "bar"
    chart5.grouping = "clustered"
    chart5.title = "Top 10 naturezas por valor"
    chart5.legend = None
    chart5.width = 16.5
    chart5.height = 7.5
    data_ref5 = Reference(ws_analises, min_col=3, min_row=29, max_row=38)
    cats_ref5 = Reference(ws_analises, min_col=1, min_row=29, max_row=38)
    chart5.add_data(data_ref5, titles_from_data=False)
    chart5.set_categories(cats_ref5)
    if chart5.series:
        apply_series_color(chart5.series[0], "2E75B6")
    ws_dash.add_chart(chart5, "A52")

    # Chart 6: BarChart (bar horizontal) - Top 10 itens por valor
    chart6 = BarChart()
    chart6.type = "bar"
    chart6.grouping = "clustered"
    chart6.title = "Top 10 itens por valor"
    chart6.legend = None
    chart6.width = 16.5
    chart6.height = 7.5
    data_ref6 = Reference(ws_analises, min_col=2, min_row=44, max_row=53)
    cats_ref6 = Reference(ws_analises, min_col=1, min_row=44, max_row=53)
    chart6.add_data(data_ref6, titles_from_data=False)
    chart6.set_categories(cats_ref6)
    if chart6.series:
        apply_series_color(chart6.series[0], "C55A11")
    ws_dash.add_chart(chart6, "N52")

    # Salvar workbook resultante
    orig_base = os.path.splitext(os.path.basename(input_path))[0]
    out_filename = f"{orig_base}_ABC_DASHBOARD.xlsx"
    if output_dir:
        output_path = os.path.join(output_dir, out_filename)
    else:
        output_path = os.path.join(os.path.dirname(input_path), out_filename)

    wb_out.save(output_path)

    top_item_val = sorted_items[0]['total_val'] if sorted_items else 0.0
    top_item_pct = (top_item_val / total_inventory_value * 100) if total_inventory_value > 0 else 0.0

    return output_path, {
        'filename': out_filename,
        'project_name': project_name,
        'date_str': date_str,
        'total_value': total_inventory_value,
        'total_value_fmt': format_currency_br(total_inventory_value),
        'distinct_products': num_products,
        'total_rows': total_base_rows,
        'locations_count': num_locs,
        'class_a_items': count_a,
        'class_a_pct_items': round((count_a / num_products * 100), 1) if num_products > 0 else 0,
        'class_a_value': val_a,
        'class_a_value_fmt': format_currency_br(val_a),
        'class_a_pct_value': round((val_a / total_inventory_value * 100), 1) if total_inventory_value > 0 else 0,
        'top_item_pct': round(top_item_pct, 1),
        'classes_summary': [
            {
                'classe': 'A',
                'itens': count_a,
                'pct_itens': f"{(count_a/num_products*100):.1f}%" if num_products > 0 else "0.0%",
                'valor_fmt': format_currency_br(val_a),
                'pct_valor': f"{(val_a/total_inventory_value*100):.1f}%" if total_inventory_value > 0 else "0.0%",
                'color': 'emerald'
            },
            {
                'classe': 'B',
                'itens': count_b,
                'pct_itens': f"{(count_b/num_products*100):.1f}%" if num_products > 0 else "0.0%",
                'valor_fmt': format_currency_br(val_b),
                'pct_valor': f"{(val_b/total_inventory_value*100):.1f}%" if total_inventory_value > 0 else "0.0%",
                'color': 'amber'
            },
            {
                'classe': 'C',
                'itens': count_c,
                'pct_itens': f"{(count_c/num_products*100):.1f}%" if num_products > 0 else "0.0%",
                'valor_fmt': format_currency_br(val_c),
                'pct_valor': f"{(val_c/total_inventory_value*100):.1f}%" if total_inventory_value > 0 else "0.0%",
                'color': 'rose'
            }
        ],
        'locations_summary': [
            {
                'codloc': loc['codloc'],
                'nome': loc['nome'],
                'itens': loc['count'],
                'valor_fmt': format_currency_br(loc['total_val']),
                'pct_valor': f"{(loc['total_val']/total_inventory_value*100):.1f}%" if total_inventory_value > 0 else "0.0%"
            } for loc in sorted_locations
        ],
        'top_5_items': [
            {
                'rank': i + 1,
                'codigo': itm['codigoprd'],
                'descricao': itm['nome'],
                'valor_fmt': format_currency_br(itm['total_val'])
            } for i, itm in enumerate(sorted_items[:5])
        ]
    }
